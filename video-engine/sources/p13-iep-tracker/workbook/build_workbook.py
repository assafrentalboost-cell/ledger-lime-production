#!/usr/bin/env python3
"""Build the Product #13 RECONSTRUCTION workbook in two states.

IMPORTANT - provenance
----------------------
The original Product #13 workbook / raw clips were not present in this
repository, so this script rebuilds the documented Product #13 logic
(goal progress, "No Recent Data" / "Improving" status, Needs Attention
count) with real spreadsheet formulas and fictional sample data. Students are
shown only as privacy-safe identifiers (STU-001...), never names.

  before : Progress Log as-is.
  after  : exactly ONE new Progress Log row is appended (the "log progress
           once" action). Nothing else changes.

Every dashboard value (60% -> 90%, No Recent Data -> Improving,
Needs Attention 2 -> 1) is computed by formulas and recalculated by
LibreOffice during capture. Nothing is typed onto the dashboard.

These captures are STAND-INS. Replace them with captures of the actual
Product #13 workbook before any Etsy use (see VIDEO-QA-REPORT-P13.md).

Usage: python build_workbook.py OUT_DIR
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

NAVY, CREAM, GOLD, INK, MUTED, LINE = "14213D", "F6F1E7", "B89B5E", "1B2433", "6B7280", "D9D2C3"
REPORT_DATE = date(2026, 9, 25)
STALE_DAYS = 30

# Goal register (fictional sample data; privacy-safe student IDs, no names).
GOALS = [
    ("G-101", "STU-001",   "Reading fluency",       "Reading", 42, 120, 120),
    ("G-102", "STU-002",   "Multi-step word problems", "Math", 35, 90, 90),
    ("G-103", "STU-003",   "Self-advocacy requests", "Behavior", 20, 60, 45),
    ("G-104", "STU-004",   "Reading comprehension", "Reading", 40, 120, 120),
    ("G-105", "STU-005",   "Written expression",    "Writing", 30, 90, 90),
    ("G-106", "STU-006",   "Math fact fluency",     "Math", 50, 60, 60),
]

# Existing Progress Log (date, goal, score %).
LOG = [
    (date(2026, 8, 28), "G-101", 55), (date(2026, 9, 18), "G-101", 64),
    (date(2026, 8, 30), "G-102", 48), (date(2026, 9, 19), "G-102", 57),
    (date(2026, 9, 2),  "G-103", 40), (date(2026, 9, 22), "G-103", 40),
    (date(2026, 7, 14), "G-104", 45), (date(2026, 8, 12), "G-104", 60),
    (date(2026, 9, 1),  "G-105", 44), (date(2026, 9, 16), "G-105", 58),
    (date(2026, 9, 3),  "G-106", 70), (date(2026, 9, 21), "G-106", 55),
]
# The single "log progress once" action.
NEW_ENTRY = (date(2026, 9, 24), "G-104", 90)

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
CREAM_FILL = PatternFill("solid", fgColor=CREAM)
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")


def page_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.3, bottom=0.3, header=0, footer=0)


def build(state: str, path: Path):
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    log = wb.create_sheet("Progress Log")
    student_of = {g[0]: g[1] for g in GOALS}

    # ---------------- Progress Log (input sheet) ----------------
    log.sheet_view.showGridLines = True
    widths = {"A": 14, "B": 14, "C": 11, "D": 12, "E": 34}
    for col, w in widths.items():
        log.column_dimensions[col].width = w
    log["A1"] = "Progress Log"
    log["A1"].font = Font(name="Liberation Sans", size=16, bold=True, color=NAVY)
    log["A2"] = "Enter one row per data collection. The dashboard updates automatically."
    log["A2"].font = Font(name="Liberation Sans", size=10, italic=True, color=MUTED)
    heads = ["Date", "Student", "Goal ID", "Score %", "Notes"]
    for i, h in enumerate(heads, 1):
        c = log.cell(row=4, column=i, value=h)
        c.font = Font(name="Liberation Sans", size=11, bold=True, color=CREAM)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center" if i < 5 else "left", vertical="center")
        c.border = BOX
    log.row_dimensions[4].height = 22
    rows = list(LOG) + ([NEW_ENTRY] if state == "after" else [])
    notes = {("G-104", 90): "Probe: 9/10 correct"}
    for r, (d, g, s) in enumerate(rows, start=5):
        vals = [d, student_of[g], g, s / 100, notes.get((g, s), "")]
        for i, v in enumerate(vals, 1):
            c = log.cell(row=r, column=i, value=v)
            c.font = Font(name="Liberation Sans", size=11, color=INK)
            c.border = BOX
            c.alignment = Alignment(horizontal="center" if i < 5 else "left", vertical="center")
        log.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        log.cell(row=r, column=4).number_format = "0%"
        log.row_dimensions[r].height = 19
    # a few empty, bordered input rows so the sheet reads as an input form
    for r in range(5 + len(rows), 5 + len(rows) + 3):
        for i in range(1, 6):
            log.cell(row=r, column=i).border = BOX
        log.row_dimensions[r].height = 19
    last_log_row = 60  # formula ranges cover future rows
    page_setup(log, landscape=False)
    log.print_area = f"A1:E{5 + len(rows) + 2}"

    # ---------------- Dashboard ----------------
    dash.sheet_view.showGridLines = False
    cols = {"A": 3, "B": 10, "C": 13, "D": 27, "E": 11, "F": 11, "G": 11, "H": 13, "I": 17, "J": 16, "K": 3}
    for col, w in cols.items():
        dash.column_dimensions[col].width = w
    for col in "LMNO":  # helper columns (outside the print area)
        dash.column_dimensions[col].width = 12
    dash.merge_cells("B2:J2")
    dash["B2"] = "IEP Caseload & Goal Progress Tracker"
    dash["B2"].font = Font(name="Liberation Serif", size=20, bold=True, color=CREAM)
    dash["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in "BCDEFGHIJ":
        dash[f"{col}2"].fill = HDR_FILL
    dash.row_dimensions[2].height = 38
    dash["B3"] = "Report date"
    dash["B3"].font = Font(name="Liberation Sans", size=10, color=MUTED)
    dash["C3"] = REPORT_DATE
    dash["C3"].number_format = "yyyy-mm-dd"
    dash["C3"].font = Font(name="Liberation Sans", size=10, color=INK)
    dash["D3"] = f"Status uses data from the last {STALE_DAYS} days"
    dash["D3"].font = Font(name="Liberation Sans", size=10, italic=True, color=MUTED)

    # KPI tiles: row 5 label, row 6 value
    first, last = 10, 10 + len(GOALS) - 1
    status_rng = f"J{first}:J{last}"
    kpis = [
        ("B5:C5", "B6:C6", "Students", f"=SUMPRODUCT(1/COUNTIF(C{first}:C{last},C{first}:C{last}))"),
        ("D5", "D6", "Active Goals", f"=COUNTA(B{first}:B{last})"),
        ("E5:G5", "E6:G6", "Needs Attention", f'=COUNTIF({status_rng},"No Recent Data")+COUNTIF({status_rng},"Declining")'),
        ("H5:J5", "H6:J6", "Improving", f'=COUNTIF({status_rng},"Improving")'),
    ]
    for lab_rng, val_rng, label, formula in kpis:
        for rng in (lab_rng, val_rng):
            if ":" in rng:
                dash.merge_cells(rng)
        lc = dash[lab_rng.split(":")[0]]
        vc = dash[val_rng.split(":")[0]]
        lc.value, vc.value = label, formula
        lc.font = Font(name="Liberation Sans", size=10, bold=True, color=MUTED)
        vc.font = Font(name="Liberation Sans", size=24, bold=True, color=NAVY)
        lc.alignment = Alignment(horizontal="center", vertical="bottom")
        vc.alignment = Alignment(horizontal="center", vertical="center")
        for rng in (lab_rng, val_rng):
            a, b = (rng.split(":") + [rng])[:2]
            for row in dash[f"{a}:{b}"]:
                for c in row:
                    c.fill = CREAM_FILL
    dash.row_dimensions[5].height = 22
    dash.row_dimensions[6].height = 40
    # gold accent under the Needs Attention tile
    for col in "EFG":
        dash[f"{col}6"].border = Border(bottom=Side(style="medium", color=GOLD))

    heads = ["Goal ID", "Student", "Goal", "Area", "Baseline", "Latest", "Last Data", "Service Min / Wk", "Status"]
    for i, h in enumerate(heads):
        c = dash.cell(row=9, column=2 + i, value=h)
        c.font = Font(name="Liberation Sans", size=10, bold=True, color=CREAM)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dash.row_dimensions[9].height = 30
    L = f"'Progress Log'!$A$5:$A${last_log_row}"
    G = f"'Progress Log'!$C$5:$C${last_log_row}"
    S = f"'Progress Log'!$D$5:$D${last_log_row}"
    for n, (gid, stu, goal, area, base, req, deliv) in enumerate(GOALS):
        r = first + n
        # helpers: L=last date, M=latest score, N=previous date, O=previous score
        dash[f"L{r}"] = f"=SUMPRODUCT(MAX(({G}=B{r})*{L}))"
        dash[f"M{r}"] = f"=SUMIFS({S},{G},B{r},{L},L{r})"
        dash[f"N{r}"] = f"=SUMPRODUCT(MAX(({G}=B{r})*({L}<L{r})*{L}))"
        dash[f"O{r}"] = f"=SUMIFS({S},{G},B{r},{L},N{r})"
        row = [gid, stu, goal, area, base / 100,
               f'=IF(L{r}=0,"",M{r})',
               f'=IF(L{r}=0,"",L{r})',
               f"{deliv} / {req}",
               f'=IF(L{r}=0,"No Data",IF($C$3-L{r}>{STALE_DAYS},"No Recent Data",'
               f'IF(M{r}>O{r},"Improving",IF(M{r}<O{r},"Declining","Steady"))))']
        for i, v in enumerate(row):
            c = dash.cell(row=r, column=2 + i, value=v)
            c.font = Font(name="Liberation Sans", size=11, color=INK, bold=(i == 8))
            c.alignment = Alignment(horizontal="left" if i in (1, 2) else "center", vertical="center")
            c.border = Border(bottom=thin)
            c.fill = WHITE_FILL
        dash[f"F{r}"].number_format = "0%"
        dash[f"G{r}"].number_format = "0%"
        dash[f"H{r}"].number_format = "yyyy-mm-dd"
        dash.row_dimensions[r].height = 24
    # status colouring (conditional formatting -> follows the recalculated value)
    rules = [("No Recent Data", "FBEAD0", "8A5A00"), ("Declining", "F8DADA", "8E1F1F"),
             ("Improving", "DCEFE0", "1E6B34"), ("Steady", "E8EAF0", "34405A")]
    for text, bg, fg in rules:
        dash.conditional_formatting.add(
            status_rng, FormulaRule(formula=[f'J{first}="{text}"'], fill=PatternFill("solid", fgColor=bg, bgColor=bg),
                                    font=Font(color=fg, bold=True)))
    page_setup(dash)
    dash.print_area = f"A1:K{last + 1}"
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    out.mkdir(parents=True, exist_ok=True)
    for st in ("before", "after"):
        p = out / f"p13_reconstruction_{st}.xlsx"
        build(st, p)
        print(p)
